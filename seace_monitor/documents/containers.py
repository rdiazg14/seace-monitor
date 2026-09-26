"""Deteccion y extraccion segura de anexos contenedores de SEACE."""
from __future__ import annotations

import io
import os
import re
import tempfile
from collections import Counter
from pathlib import Path

import pymupdf

from seace_monitor.documents.pdf_extraction import chars_utiles, limpiar_texto

MAX_ARCHIVOS_ZIP = 100
MAX_BYTES_ZIP = 50 * 1024 * 1024
MAX_IMG_OCR = 5
MIN_CHARS_UTIL = 200
_TZ_NOMBRE_IGNORABLE = re.compile(
    r"(?i)(^|/)(__MACOSX|\.DS_Store|Thumbs\.db|~\$.*\.doc[x]?)$"
)
_EXTS_IMG = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}
MOTIVO_RAR = "contenedor rar sin binario (unrar/7z)"
MOTIVO_DOC = "contenedor doc binario (OLE2) sin soporte"
MOTIVO_DESC = "contenedor no procesable"

def _head(body: bytes, n: int = 8) -> bytes:
    return (body or b"")[:n]


def _es_zip(body: bytes) -> bool:
    h = _head(body, 4)
    return h == b"PK\x03\x04" or h == b"PK\x05\x06" or h == b"PK\x07\x08"


def _es_rar(body: bytes) -> bool:
    h = _head(body, 8)
    return h.startswith(b"Rar!\x1a\x07\x00") or h.startswith(b"Rar!\x1a\x07\x01")


def _es_ole2(body: bytes) -> bool:
    return _head(body, 8) == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _zip_kind(body: bytes) -> str:
    """Distingue docx/xlsx/pptx/zip genérico por su estructura interna."""
    import zipfile

    try:
        with zipfile.ZipFile(io.BytesIO(body)) as zf:
            names = set(zf.namelist())
    except Exception:
        return "zip"
    if "word/document.xml" in names:
        return "docx"
    if "xl/workbook.xml" in names:
        return "xlsx"
    if "ppt/presentation.xml" in names:
        return "pptx"
    return "zip"


def tipo_contenedor(body: bytes, nombre: str = "") -> str:
    """Tipo real por magic bytes (el mime/extensión de SEACE NO es fiable)."""
    b = body or b""
    if not b:
        return "vacio"
    if b.lstrip().startswith(b"%PDF"):
        return "pdf"
    if _es_rar(b):
        return "rar"
    if _es_ole2(b):
        return "ole2"
    if _es_zip(b):
        return _zip_kind(b)
    ext = (nombre or "").lower().rsplit(".", 1)[-1] if "." in (nombre or "") else ""
    if ext in _EXTS_IMG:
        return "imagen"
    if b[:4] == b"\x89PNG":
        return "imagen"
    if b[:2] == b"\xff\xd8":
        return "imagen"
    return "desconocido"


def _mime_imagen(body: bytes) -> str:
    if body[:4] == b"\x89PNG":
        return "image/png"
    if body[:2] == b"\xff\xd8":
        return "image/jpeg"
    return "image/jpeg"


# ── Extractores por formato ───────────────────────────────────────────────────
def _docx_texto_nativo(body: bytes) -> str:
    """Texto de word/document.xml: TODO `w:p` en orden (incluye tablas y cajas)."""
    import zipfile

    from lxml import etree

    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    with zipfile.ZipFile(io.BytesIO(body)) as zf:
        xml = zf.read("word/document.xml")
    root = etree.fromstring(xml)
    lineas: list[str] = []
    for p in root.iter(f"{{{W}}}p"):
        trozos = [t.text or "" for t in p.iter(f"{{{W}}}t")]
        linea = "".join(trozos).strip()
        if linea:
            lineas.append(linea)
    return limpiar_texto("\n".join(lineas))


def _docx_imagenes(body: bytes) -> list[bytes]:
    """Imágenes embebidas (word/media/*) en orden numérico. Escaneos DOCX."""
    import zipfile

    with zipfile.ZipFile(io.BytesIO(body)) as zf:
        nombres = [n for n in zf.namelist() if n.lower().startswith("word/media/")]
        nombres.sort(key=lambda n: _num_img(n))
        out: list[bytes] = []
        for n in nombres:
            data = zf.read(n)
            if data:
                out.append(data)
    return out


def _num_img(nombre: str) -> int:
    m = re.search(r"(\d+)", nombre.rsplit("/", 1)[-1])
    return int(m.group(1)) if m else 0


def extraer_docx(
    body: bytes, *, ocr_page=None
) -> tuple[str, dict]:
    """Texto de un DOCX. Si es escaneado (sin texto nativo, con imágenes),
    hace OCR de las imágenes embebidas (tope MAX_IMG_OCR)."""
    stats: dict = {"nativo_chars": 0, "imagenes": 0, "ocr_img": 0}
    nativo = _docx_texto_nativo(body)
    stats["nativo_chars"] = chars_utiles(nativo)
    if stats["nativo_chars"] >= MIN_CHARS_UTIL:
        return nativo, stats
    imgs = _docx_imagenes(body)
    stats["imagenes"] = len(imgs)
    if not imgs or ocr_page is None:
        return nativo, stats
    bloques = [nativo] if nativo.strip() else []
    for i, img in enumerate(imgs, 1):
        if stats["ocr_img"] >= MAX_IMG_OCR:
            break
        try:
            ocr = ocr_page(img, _mime_imagen(img))
            bloques.append(f"--- pagina {i} (ocr) ---\n{ocr.strip()}")
            stats["ocr_img"] += 1
        except Exception:
            continue
    return limpiar_texto("\n\n".join(bloques)), stats


def extraer_xlsx(body: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    out: list[str] = []
    for ws in wb.worksheets:
        out.append(f"--- hoja {ws.title} ---")
        for row in ws.iter_rows(values_only=True):
            celdas = ["" if v is None else str(v).strip() for v in row]
            if any(celdas):
                out.append(" | ".join(celdas))
    wb.close()
    return limpiar_texto("\n".join(out))


def extraer_pdf_bytes(body: bytes) -> str:
    with pymupdf.open(stream=body, filetype="pdf") as doc:
        paginas = [(page.get_text("text") or "").strip() for page in doc]
    return limpiar_texto("\n".join(p for p in paginas if p))


def _nombre_ignorable(nombre: str) -> bool:
    base = nombre.replace("\\", "/").split("/")[-1]
    if base in (".DS_Store", "Thumbs.db"):
        return True
    if base.startswith("~$"):
        return True
    return bool(_TZ_NOMBRE_IGNORABLE.search(nombre))


def extraer_zip(
    body: bytes, *, ocr_page=None
) -> tuple[str, dict]:
    """Extrae texto de todos los hijos procesables de un ZIP (en memoria).

    Devuelve (texto, stats). No escribe a disco: inmune a path traversal.
    """
    import zipfile

    stats: dict = {"archivos": 0, "por_tipo": Counter(), "ocr_img": 0, "omitidos": 0}
    try:
        zf = zipfile.ZipFile(io.BytesIO(body))
    except Exception as e:
        raise RuntimeError(f"zip ilegible: {e}") from e

    with zf:
        infos = [i for i in zf.infolist() if not i.is_dir()]
        if len(infos) > MAX_ARCHIVOS_ZIP:
            raise RuntimeError(f"zip con {len(infos)} archivos (> {MAX_ARCHIVOS_ZIP})")
        total_bytes = sum(i.file_size for i in infos)
        if total_bytes > MAX_BYTES_ZIP:
            raise RuntimeError(
                f"zip descomprimido {total_bytes} bytes (> {MAX_BYTES_ZIP})"
            )
        infos.sort(key=lambda i: i.filename.lower())
        bloques: list[str] = []
        stats["archivos"] = len(infos)
        for info in infos:
            nombre = info.filename
            if _nombre_ignorable(nombre):
                stats["omitidos"] += 1
                continue
            try:
                data = zf.read(info)
            except Exception:
                stats["omitidos"] += 1
                continue
            if not data:
                continue
            tipo = tipo_contenedor(data, nombre)
            stats["por_tipo"][tipo] += 1
            texto = ""
            try:
                if tipo == "pdf":
                    texto = extraer_pdf_bytes(data)
                elif tipo == "docx":
                    texto, _s = extraer_docx(data, ocr_page=ocr_page)
                elif tipo == "xlsx":
                    texto = extraer_xlsx(data)
                elif tipo == "imagen":
                    if ocr_page is not None and stats["ocr_img"] < MAX_IMG_OCR:
                        texto = ocr_page(data, _mime_imagen(data))
                        stats["ocr_img"] += 1
                # rar anidado / ole2 / pptx / desconocido: se omiten (silencioso)
            except Exception:
                texto = ""
            if (texto or "").strip():
                bloques.append(f"--- {nombre} ---\n{texto.strip()}")
    texto = limpiar_texto("\n\n".join(bloques))
    return texto, stats


_RAR_TOOLS = ("unrar", "unrar-free", "7z", "7za", "7zz", "bsdtar", "unar")


def _find_rar_tool() -> str | None:
    import shutil

    for t in _RAR_TOOLS:
        if shutil.which(t):
            return t
    return None


def extraer_rar(
    body: bytes, *, ocr_page=None
) -> tuple[str, dict]:
    """Extrae texto de un RAR vía rarfile (requiere binario unrar/7z/bsdtar).

    Lanza RuntimeError si no hay backend. El caller decide si registrar rechazo.
    """
    import rarfile

    tool = _find_rar_tool()
    if not tool:
        raise RuntimeError(MOTIVO_RAR)
    rarfile.UNRAR_TOOL = tool
    fd, tmp_name = tempfile.mkstemp(prefix="seace-rar-", suffix=".rar")
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        tmp.write_bytes(body)
        rf = rarfile.RarFile(str(tmp))
        infos = [i for i in rf.infolist() if not i.is_dir()]
        if len(infos) > MAX_ARCHIVOS_ZIP:
            raise RuntimeError(f"rar con {len(infos)} archivos (> {MAX_ARCHIVOS_ZIP})")
        bloques: list[str] = []
        stats: dict = {"archivos": len(infos), "por_tipo": Counter(), "ocr_img": 0}
        for info in sorted(infos, key=lambda i: i.filename.lower()):
            if _nombre_ignorable(info.filename):
                continue
            data = rf.read(info.filename)
            if not data:
                continue
            tipo = tipo_contenedor(data, info.filename)
            stats["por_tipo"][tipo] += 1
            texto = ""
            try:
                if tipo == "pdf":
                    texto = extraer_pdf_bytes(data)
                elif tipo == "docx":
                    texto, _s = extraer_docx(data, ocr_page=ocr_page)
                elif tipo == "xlsx":
                    texto = extraer_xlsx(data)
                elif tipo == "imagen":
                    if ocr_page is not None and stats["ocr_img"] < MAX_IMG_OCR:
                        texto = ocr_page(data, _mime_imagen(data))
                        stats["ocr_img"] += 1
            except Exception:
                texto = ""
            if (texto or "").strip():
                bloques.append(f"--- {info.filename} ---\n{texto.strip()}")
        return limpiar_texto("\n\n".join(bloques)), stats
    finally:
        tmp.unlink(missing_ok=True)


# ── Elección de anexo contenedor ──────────────────────────────────────────────
_PRIORIDAD_EXT = {"docx": 0, "zip": 1, "rar": 2, "doc": 3}


def _ext(nombre: str) -> str:
    n = (nombre or "").lower().strip()
    return n.rsplit(".", 1)[-1] if "." in n else ""


def elegir_contenedor(archivos: list) -> dict | None:
    """Elige el mejor anexo NO-PDF. Prefiere docx, luego zip, rar, doc.

    No compite con `elegir_pdf`: este script corre sobre contratos donde ya
    se descartó el PDF (o no hay). Devuelve None si no hay contenedor.
    """
    cands: list[tuple[int, dict]] = []
    for a in archivos:
        if not isinstance(a, dict):
            continue
        nombre = str(a.get("nombre") or "")
        mime = str(a.get("descripcionMime") or "").lower()
        ext = _ext(nombre)
        if ext in _PRIORIDAD_EXT:
            cands.append((_PRIORIDAD_EXT[ext], a))
        elif "word" in mime or "msword" in mime:
            cands.append((0, a))
        elif "zip" in mime or "zip" in nombre.lower():
            cands.append((1, a))
        elif "rar" in mime or "rar" in nombre.lower():
            cands.append((2, a))
    if not cands:
        return None
    cands.sort(key=lambda x: x[0])
    return cands[0][1]


# ── Descarga cruda (sin validación PDF) ───────────────────────────────────────
def _parece_html(body: bytes) -> bool:
    sample = (body or b"")[:400].lstrip().lower()
    return (
        sample.startswith(b"<!doctype")
        or sample.startswith(b"<html")
        or b"<html" in sample[:200]
    )


def descargar_crudo(http: SeaceHttp, url: str) -> bytes:
    status, headers, body = http.get_bytes(url)
    ctype = headers.get("content-type") or ""
    if status != 200:
        raise RuntimeError(f"descargar HTTP {status} ({ctype[:80]})")
    if not body:
        raise RuntimeError("respuesta vacía al descargar contenedor")
    if _parece_html(body) or "json" in ctype.lower() or "text/html" in ctype.lower():
        raise RuntimeError(f"binario parece HTML/JSON (content-type={ctype[:80]})")
    return body


# ── Selección de vigentes TI/IA sin TDR ───────────────────────────────────────
